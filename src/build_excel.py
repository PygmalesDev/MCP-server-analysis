"""Collect all analysis results (data/<prompt version>/*.json) into one Excel workbook.

Settings: config.py (EXCEL_DATA_DIRS = which prompt versions; a later version replaces an earlier result). Run: python main.py --excel
Output: data/analysis_result/MCP_analysis_results.xlsx (close it in Excel before running again).
"""

import csv, json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

OUT = config.RESULT_DIR / "MCP_analysis_results.xlsx"
CAPS = ["Tools", "Resources", "Prompts", "Logging", "Sampling", "Roots", "Authorization", "Elicitation",
        "Tasks", "Configuration", "Error Reporting", "Subscription", "Progress Tracking", "Cancellation", "Other"]
COLORS = {"meta": "404040", "A": "1F4E78", "B": "7F6000", "C": "375623", "D": "833C0B", "E": "3A3838", "F": "512E5F"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def text(value) -> str:
    """Lists/dicts -> readable text; Excel allows at most 32,767 characters per cell."""
    if isinstance(value, list):
        value = "\n".join(text(v) for v in value)
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False)
    value = "" if value is None else str(value)
    return value if len(value) < 32000 else value[:32000] + " …[truncated]"


def sheet(wb, title, columns, rows, first=False):
    """columns: [(group, header, width)]; rows: list of value lists."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for i, (group, header, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=COLORS[group])
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 48
    for r, values in enumerate(rows, 2):
        for i, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=i, value=value if isinstance(value, (int, float)) else text(value))
            cell.font = Font(name="Arial", size=10, bold=(i == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
        ws.row_dimensions[r].height = 60
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 2)}"


def all_sources(result: dict) -> list:
    found = []
    for section in result.values():
        if isinstance(section, dict):
            found += section.get("sources", [])
    for cap in result.get("C_capabilities", []):
        found += cap.get("sources", [])
    return list(dict.fromkeys(str(s) for s in found))


def main():
    with open(config.CSV_FILE, encoding="utf-8") as f:
        csv_rows = {row["name"]: row for row in csv.DictReader(f)}
    results = {}
    for folder in getattr(config, "EXCEL_DATA_DIRS", [config.DATA_DIR]):   # later folders replace earlier results
        for p in sorted(folder.glob("*.json")):
            results[p.stem] = json.loads(p.read_text(encoding="utf-8"))
    results = dict(sorted(results.items()))
    manual = set(config.CHOSEN_FILE.read_text(encoding="utf-8").split())
    if not results:
        raise SystemExit(f"No results in {getattr(config, 'EXCEL_DATA_DIRS', [config.DATA_DIR])}")

    servers, caps, actions, observations = [], [], [], []
    for name, d in results.items():
        A, B = d.get("A_general", {}), d.get("B_access_pricing", {})
        D, E, F = d.get("D_server_side_logic", {}), d.get("E_aggregated", {}), d.get("F_documentation", {})
        run, row = d.get("_run", {}), csv_rows.get(name, {})
        acts = E.get("actions") if isinstance(E.get("actions"), list) else []
        cap_list = d.get("C_capabilities", [])
        implemented = sum(1 for c in cap_list if str(c.get("implemented", "")).startswith(("Yes", "Partially")))
        servers.append([
            name, "Done (LLM draft)", "Yes" if name in manual else "No",
            f"{d.get('analyst_method', '')}; {run.get('model', '')}, effort {run.get('reasoning_effort', '')}",
            run.get("date", d.get("analysis_date", "")),
            run.get("duration_min") if isinstance(run.get("duration_min"), (int, float)) else "",
            row.get("endpoint", A.get("endpoint")), row.get("docs", A.get("documentation_url_csv")),
            row.get("addressing", A.get("addressing")), A.get("repository_url"), A.get("documentation_urls_used"),
            f"{A.get('open_code_equals_deployed_server', '')}\n{A.get('open_code_note', '')}".strip(),
            A.get("mcp_protocol_version"), A.get("commit_hash"), A.get("commit_date"),
            A.get("implementation_language"), A.get("mcp_sdk_and_version"), A.get("code_public"), A.get("license"),
            A.get("release_status"), A.get("other_official_servers"),
            B.get("mcp_access_cost"), B.get("mcp_access_cost_details"), B.get("linked_software"),
            B.get("linked_software_cost"), B.get("pricing_model_details"), B.get("other_access_restrictions"),
            B.get("restrictions_controlled_by"),
            implemented,
            D.get("additional_logic"), D.get("behavioural_instructions_to_ai_verbatim"),
            D.get("server_structure"), D.get("server_structure_explanation"),
            len(acts), f"{E.get('action_count', '')}\n{E.get('action_count_comment', '')}".strip(),
            E.get("pre_interaction_tokens"), E.get("token_method"), E.get("human_can_agent_cannot"),
            F.get("implementation_hints_not_enforced"), F.get("governance_promises"), F.get("promises_to_user"),
            F.get("comparison_with_server_logic"),
            d.get("confidence"), all_sources(d), d.get("notes"),
        ])
        for c in cap_list:
            caps.append([name, c.get("capability"), c.get("implemented"), c.get("depends_on_mcp_version"),
                         c.get("description_what_how"), c.get("pre_interaction_context_verbatim"),
                         c.get("post_interaction_context_structure"), c.get("authorization_mechanism", ""),
                         c.get("authorization_server", ""), c.get("sources")])
        for i, a in enumerate(acts, 1):
            actions.append([name, i, a.get("tool_name"), a.get("description"), a.get("type"),
                            a.get("enabled_by_default"), a.get("parameters"), a.get("parameter_count"),
                            a.get("source"), a.get("source_type")])
        observations += [[name, "Contradiction", x] for x in d.get("contradictions", [])]
        observations += [[name, "Observation for research", x] for x in d.get("observations_for_research", [])]

    wb = Workbook()
    sheet(wb, "Servers", [
        ("meta", "Server", 16), ("meta", "Status", 12), ("meta", "Also analysed manually?", 11), ("meta", "Analyst / method", 22), ("meta", "Analysis date", 11),
        ("meta", "Duration, min", 9),
        ("A", "Endpoint (from CSV)", 28), ("A", "Documentation (from CSV)", 28), ("A", "Addressing", 9),
        ("A", "Code repository", 28), ("A", "Documentation used", 36), ("A", "Open code = deployed server?", 30),
        ("A", "MCP version", 14), ("A", "Commit hash", 22), ("A", "Commit date", 11), ("A", "Language", 12),
        ("A", "MCP SDK and version", 20), ("A", "Code public?", 10), ("A", "License", 12), ("A", "Release status", 14),
        ("A", "Other official servers", 36),
        ("B", "MCP access: free / paid", 12), ("B", "MCP access cost", 30), ("B", "Linked software", 14),
        ("B", "Linked software cost", 36), ("B", "Pricing model", 40), ("B", "Other access restrictions", 40),
        ("B", "Restrictions controlled by", 28),
        ("C", "Capabilities implemented (Yes/Partially)", 13),
        ("D", "Additional server-side logic (checklist 1-13)", 70), ("D", "Behavioural instructions to the AI (verbatim)", 60),
        ("D", "Server structure", 13), ("D", "Structure: explanation", 36),
        ("E", "Number of actions (sheet Actions)", 11), ("E", "Action count comment", 40),
        ("E", "Tokens before first action", 12), ("E", "Token counting method", 30),
        ("E", "Human can, agent cannot", 50),
        ("F", "Recommendations not enforced by the server", 50), ("F", "Governance promises", 50),
        ("F", "Promises to the user", 40), ("F", "Promises vs. server logic", 40),
        ("meta", "Confidence", 16), ("meta", "Sources", 50), ("meta", "Notes", 50),
    ], servers, first=True)
    sheet(wb, "Capabilities", [
        ("meta", "Server", 16), ("C", "Capability", 16), ("C", "Implemented?", 14), ("C", "Depends on MCP version", 16),
        ("C", "Description (what / how)", 50), ("C", "Pre-Interaction Context (verbatim)", 60),
        ("C", "Post-Interaction Context (structure)", 50), ("C", "Authorization: mechanism", 40),
        ("C", "Authorization: server", 24), ("C", "Sources", 40),
    ], caps)
    sheet(wb, "Actions", [
        ("meta", "Server", 16), ("E", "#", 6), ("E", "Tool (name in MCP)", 28), ("E", "Description", 44),
        ("E", "Type", 10), ("E", "Enabled by default?", 14), ("E", "Parameters", 60), ("E", "Number of parameters", 10),
        ("E", "Source", 36), ("E", "Source type", 24),
    ], actions)
    sheet(wb, "Observations", [
        ("meta", "Server", 16), ("F", "Type", 22), ("F", "Text", 100),
    ], observations)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT}: {len(servers)} servers, {len(actions)} actions, {len(observations)} observations")
